from __future__ import annotations

import json
from dataclasses import dataclass
from datetime import date
from pathlib import Path
from typing import Iterable

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


@dataclass(frozen=True)
class Theme:
    font_name: str = "Calibri"
    text_light: RGBColor = RGBColor(255, 255, 255)
    text_dark: RGBColor = RGBColor(20, 24, 31)
    accent: RGBColor = RGBColor(104, 151, 255)


def _make_linear_gradient_png(
    path: Path,
    size_px: tuple[int, int],
    left_rgb: tuple[int, int, int],
    right_rgb: tuple[int, int, int],
) -> Path:
    w, h = size_px
    base = Image.new("RGB", (w, h), left_rgb)
    top = Image.new("RGB", (w, h), right_rgb)
    mask = Image.new("L", (w, 1))
    for x in range(w):
        mask.putpixel((x, 0), int(255 * (x / max(1, w - 1))))
    mask = mask.resize((w, h))
    out = Image.composite(top, base, mask)
    out.save(path, format="PNG", optimize=True)
    return path


def _add_full_slide_bg(slide, bg_path: Path, prs: Presentation) -> None:
    slide.shapes.add_picture(str(bg_path), 0, 0, width=prs.slide_width, height=prs.slide_height)


def _add_title(slide, text: str, theme: Theme, x, y, w, h, size_pt: int = 40) -> None:
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = text
    run.font.name = theme.font_name
    run.font.size = Pt(size_pt)
    run.font.bold = True
    run.font.color.rgb = theme.text_light
    p.alignment = PP_ALIGN.LEFT


def _add_subtitle(slide, text: str, theme: Theme, x, y, w, h, size_pt: int = 18) -> None:
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = text
    run.font.name = theme.font_name
    run.font.size = Pt(size_pt)
    run.font.color.rgb = theme.text_light
    p.alignment = PP_ALIGN.LEFT


def _add_bullets(slide, title: str, bullets: list[str], theme: Theme, x, y, w, h) -> None:
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.clear()

    p0 = tf.paragraphs[0]
    p0.text = title
    p0.font.name = theme.font_name
    p0.font.size = Pt(24)
    p0.font.bold = True
    p0.font.color.rgb = theme.text_light

    for b in bullets:
        p = tf.add_paragraph()
        p.text = b
        p.level = 0
        p.font.name = theme.font_name
        p.font.size = Pt(16)
        p.font.color.rgb = theme.text_light


def _add_image(slide, img_path: Path, x, y, w=None, h=None, border: bool = True) -> None:
    pic = slide.shapes.add_picture(str(img_path), x, y, width=w, height=h)
    if border:
        line = pic.line
        line.color.rgb = RGBColor(255, 255, 255)
        line.width = Pt(1)


def _load_metrics(metrics_json_path: Path) -> dict:
    return json.loads(metrics_json_path.read_text(encoding="utf-8"))


def _best_model_row(metrics: dict) -> dict:
    rows = metrics["model_comparison"]
    return sorted(rows, key=lambda r: (r["f1_macro"], r["accuracy"]), reverse=True)[0]


def build_ppt(
    pptx_path: Path,
    plots: Iterable[Path],
    metrics_json_path: Path,
    name_placeholder: str = "Name: ____________________",
) -> Path:
    pptx_path = Path(pptx_path)
    pptx_path.parent.mkdir(parents=True, exist_ok=True)

    theme = Theme()
    prs = Presentation()
    prs.slide_width = Inches(13.33)
    prs.slide_height = Inches(7.5)

    metrics = _load_metrics(metrics_json_path)
    best = _best_model_row(metrics)

    plots = list(plots)
    pairplot = next((p for p in plots if p.name.startswith("01_")), None)
    heatmap = next((p for p in plots if p.name.startswith("02_")), None)
    dists = next((p for p in plots if p.name.startswith("03_")), None)
    table = next((p for p in plots if p.name.startswith("04_")), None)

    bg_dir = pptx_path.parent
    bg1 = _make_linear_gradient_png(bg_dir / "_bg1.png", (1600, 900), (15, 22, 38), (84, 36, 196))
    bg2 = _make_linear_gradient_png(bg_dir / "_bg2.png", (1600, 900), (12, 32, 60), (30, 140, 150))
    bg3 = _make_linear_gradient_png(bg_dir / "_bg3.png", (1600, 900), (20, 20, 20), (70, 95, 160))
    bg4 = _make_linear_gradient_png(bg_dir / "_bg4.png", (1600, 900), (16, 26, 40), (120, 92, 255))
    bg5 = _make_linear_gradient_png(bg_dir / "_bg5.png", (1600, 900), (10, 25, 35), (10, 80, 120))

    blank = prs.slide_layouts[6]

    # Slide 1: Title
    s1 = prs.slides.add_slide(blank)
    _add_full_slide_bg(s1, bg1, prs)
    _add_title(
        s1,
        "Iris Flower Species\nClassification",
        theme,
        x=Inches(0.9),
        y=Inches(1.4),
        w=Inches(11.5),
        h=Inches(2.2),
        size_pt=54,
    )
    _add_subtitle(
        s1,
        f"{name_placeholder}\nDate: {date.today().isoformat()}",
        theme,
        x=Inches(0.95),
        y=Inches(4.2),
        w=Inches(11.5),
        h=Inches(1.0),
        size_pt=18,
    )

    # Slide 2: Problem + Dataset
    s2 = prs.slides.add_slide(blank)
    _add_full_slide_bg(s2, bg2, prs)
    _add_bullets(
        s2,
        "Problem",
        [
            "Predict Iris species from 4 flower measurements.",
            "A classic multi-class classification task (3 classes).",
        ],
        theme,
        x=Inches(0.8),
        y=Inches(0.7),
        w=Inches(6.0),
        h=Inches(3.0),
    )
    _add_bullets(
        s2,
        "Dataset",
        [
            "Source: sklearn Iris dataset (150 rows).",
            "Features: sepal length/width, petal length/width.",
            "Classes: setosa, versicolor, virginica.",
        ],
        theme,
        x=Inches(0.8),
        y=Inches(3.6),
        w=Inches(6.2),
        h=Inches(3.4),
    )
    if pairplot:
        _add_image(s2, pairplot, x=Inches(7.2), y=Inches(1.1), w=Inches(5.6), h=Inches(5.8))

    # Slide 3: EDA + Preprocessing
    s3 = prs.slides.add_slide(blank)
    _add_full_slide_bg(s3, bg3, prs)
    _add_bullets(
        s3,
        "EDA + Preprocessing",
        [
            "Validated numeric features, removed duplicates, ensured clean labels.",
            "All 4 standard features retained (high signal on petal measures).",
            "Train-test split with stratification for fair evaluation.",
        ],
        theme,
        x=Inches(0.8),
        y=Inches(0.7),
        w=Inches(12.2),
        h=Inches(1.8),
    )
    if heatmap:
        _add_image(s3, heatmap, x=Inches(0.9), y=Inches(2.6), w=Inches(6.0), h=Inches(4.6))
    if dists:
        _add_image(s3, dists, x=Inches(7.0), y=Inches(2.6), w=Inches(5.8), h=Inches(4.6))

    # Slide 4: Models + Metrics
    s4 = prs.slides.add_slide(blank)
    _add_full_slide_bg(s4, bg4, prs)
    _add_bullets(
        s4,
        "Modeling + Metrics",
        [
            "Models: Logistic Regression, Decision Tree, Random Forest.",
            "Compared with Accuracy / Precision / Recall / F1 (macro).",
            f"Best model: {best['name']} (F1={best['f1_macro']:.3f}, Acc={best['accuracy']:.3f}).",
        ],
        theme,
        x=Inches(0.8),
        y=Inches(0.7),
        w=Inches(12.2),
        h=Inches(1.8),
    )
    if table:
        _add_image(s4, table, x=Inches(0.9), y=Inches(2.5), w=Inches(12.0), h=Inches(2.1))
    if pairplot:
        _add_image(s4, pairplot, x=Inches(0.9), y=Inches(4.9), w=Inches(5.7), h=Inches(2.4))
    if heatmap:
        _add_image(s4, heatmap, x=Inches(7.2), y=Inches(4.9), w=Inches(5.7), h=Inches(2.4))

    # Slide 5: Conclusion
    s5 = prs.slides.add_slide(blank)
    _add_full_slide_bg(s5, bg5, prs)
    _add_title(
        s5,
        "Conclusion",
        theme,
        x=Inches(0.9),
        y=Inches(0.8),
        w=Inches(11.5),
        h=Inches(1.0),
        size_pt=44,
    )
    _add_bullets(
        s5,
        "Key takeaways",
        [
            f"Achieved >90% accuracy (best test accuracy: {best['accuracy']:.3f}).",
            "Petal measurements are the strongest separators between classes.",
            "Deployed the trained model via an interactive Streamlit app.",
        ],
        theme,
        x=Inches(0.9),
        y=Inches(2.0),
        w=Inches(12.0),
        h=Inches(2.4),
    )
    _add_bullets(
        s5,
        "Future scope",
        [
            "Add input validation + model confidence/probabilities.",
            "Track experiments with MLflow and add CI checks.",
            "Extend to other flower datasets or add SHAP explanations.",
        ],
        theme,
        x=Inches(0.9),
        y=Inches(4.7),
        w=Inches(12.0),
        h=Inches(2.5),
    )

    prs.save(str(pptx_path))
    return pptx_path

